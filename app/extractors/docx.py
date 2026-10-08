from pathlib import Path

from docx import Document

from .base import DocumentExtractor, ExtractedDocument, ExtractionError


class DocxExtractor(DocumentExtractor):
    extensions = frozenset({".docx"})

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            document = Document(path)
        except Exception as exc:
            raise ExtractionError("The Word document is corrupted or could not be read.") from exc

        blocks = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))

        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        font_sizes = [
            run.font.size.pt for paragraph in paragraphs for run in paragraph.runs
            if run.text.strip() and run.font.size is not None
        ]
        font_names = [
            run.font.name for paragraph in paragraphs for run in paragraph.runs
            if run.text.strip() and run.font.name
        ]
        line_spacings = [
            paragraph.paragraph_format.line_spacing for paragraph in paragraphs
            if isinstance(paragraph.paragraph_format.line_spacing, (int, float))
        ]
        section_margins = [
            {
                "top": round(section.top_margin.inches, 2) if section.top_margin else None,
                "bottom": round(section.bottom_margin.inches, 2) if section.bottom_margin else None,
                "left": round(section.left_margin.inches, 2) if section.left_margin else None,
                "right": round(section.right_margin.inches, 2) if section.right_margin else None,
            }
            for section in document.sections
        ]
        headings = sum(paragraph.style.name.lower().startswith("heading") for paragraph in paragraphs)
        formatting = {
            "font_names": font_names,
            "font_sizes": font_sizes,
            "average_font_size": round(sum(font_sizes) / len(font_sizes), 1) if font_sizes else None,
            "line_spacing": line_spacings,
            "margins": section_margins,
            "heading_count": headings,
            "paragraph_count": len(paragraphs),
        }
        return ExtractedDocument("\n".join(blocks), formatting)
