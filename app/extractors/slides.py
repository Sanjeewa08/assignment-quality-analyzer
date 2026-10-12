import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation

from .base import DocumentExtractor, ExtractedDocument, ExtractionError


def legacy_powerpoint_converter():
    return shutil.which("soffice") or shutil.which("libreoffice")


class PowerPointExtractor(DocumentExtractor):
    extensions = frozenset({".pptx"})

    def extract(self, path: Path) -> ExtractedDocument:
        try:
            presentation = Presentation(path)
        except Exception as exc:
            raise ExtractionError("The PowerPoint file is corrupted or could not be read.") from exc
        slide_text = []
        for number, slide in enumerate(presentation.slides, start=1):
            items = [shape.text.strip() for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            if items:
                slide_text.append(f"Slide {number}\n" + "\n".join(items))
        return ExtractedDocument("\n\n".join(slide_text))


class LegacyPowerPointExtractor(DocumentExtractor):
    extensions = frozenset({".ppt"})

    def extract(self, path: Path) -> ExtractedDocument:
        converter = legacy_powerpoint_converter()
        if not converter:
            raise ExtractionError("Legacy .ppt files require LibreOffice to be installed for safe conversion.")
        try:
            with tempfile.TemporaryDirectory(prefix="aqa-ppt-") as output_directory:
                subprocess.run(
                    [
                        converter,
                        "--headless",
                        "--convert-to",
                        "pptx",
                        "--outdir",
                        output_directory,
                        str(path),
                    ],
                    check=True,
                    capture_output=True,
                    timeout=60,
                )
                converted = Path(output_directory) / f"{path.stem}.pptx"
                if not converted.is_file():
                    raise ExtractionError("LibreOffice could not convert this .ppt presentation.")
                return PowerPointExtractor().extract(converted)
        except (OSError, subprocess.SubprocessError) as exc:
            raise ExtractionError("LibreOffice could not convert this .ppt presentation.") from exc
