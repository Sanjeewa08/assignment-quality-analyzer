from io import BytesIO

import pytest
from werkzeug.datastructures import FileStorage

from app.extractors import ExtractionError, extract_text


def storage(name, content):
    return FileStorage(stream=BytesIO(content), filename=name)


def test_txt_extraction():
    assert extract_text(storage("essay.txt", b"Hello world")) == "Hello world"


def test_unsupported_format_is_rejected():
    with pytest.raises(ExtractionError, match="Unsupported|Supported"):
        extract_text(storage("essay.exe", b"no"))


def test_corrupt_docx_is_rejected():
    with pytest.raises(ExtractionError, match="corrupted"):
        extract_text(storage("essay.docx", b"not a zip archive"))


def test_empty_filename_is_rejected():
    with pytest.raises(ExtractionError, match="choose"):
        extract_text(storage("", b""))


def test_extractor_factory_routes_supported_extensions():
    from pathlib import Path

    from app.extractors.factory import ExtractorFactory
    from app.extractors.docx import DocxExtractor
    from app.extractors.image import ImageExtractor
    from app.extractors.pdf import PdfExtractor
    from app.extractors.slides import LegacyPowerPointExtractor, PowerPointExtractor
    from app.extractors.text import TextExtractor

    classes = {
        ".txt": TextExtractor,
        ".docx": DocxExtractor,
        ".pdf": PdfExtractor,
        ".pptx": PowerPointExtractor,
        ".ppt": LegacyPowerPointExtractor,
        ".jpg": ImageExtractor,
        ".jpeg": ImageExtractor,
        ".png": ImageExtractor,
    }
    for extension, expected in classes.items():
        assert isinstance(ExtractorFactory.for_path(Path("assignment" + extension)), expected)


def test_factory_rejects_unsupported_extension():
    from pathlib import Path

    from app.extractors.factory import ExtractorFactory

    with pytest.raises(ExtractionError, match="Supported formats"):
        ExtractorFactory.for_path(Path("bad.exe"))


def test_scanned_pdf_reports_missing_ocr_dependencies(monkeypatch, tmp_path):
    from reportlab.pdfgen import canvas

    from app.extractors.pdf import PdfExtractor

    path = tmp_path / "scanned.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.save()
    monkeypatch.setattr(
        PdfExtractor,
        "_ocr",
        staticmethod(lambda _path: (_ for _ in ()).throw(ExtractionError("No text was found. Scanned PDFs need Poppler and Tesseract OCR installed."))),
    )
    with pytest.raises(ExtractionError, match="Tesseract"):
        PdfExtractor().extract(path)


def test_pptx_extractor_reads_slide_text(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from app.extractors.slides import PowerPointExtractor

    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    textbox.text = "Research overview"
    presentation.save(path)
    assert "Research overview" in PowerPointExtractor().extract(path).text


def test_legacy_ppt_fails_with_install_guidance_when_libreoffice_missing(monkeypatch, tmp_path):
    from app.extractors.slides import LegacyPowerPointExtractor

    monkeypatch.setattr("app.extractors.slides.shutil.which", lambda _name: None)
    path = tmp_path / "old.ppt"
    path.write_bytes(b"legacy")
    with pytest.raises(ExtractionError, match="LibreOffice"):
        LegacyPowerPointExtractor().extract(path)


def test_docx_extractor_reads_content_and_formatting(tmp_path):
    from docx import Document

    from app.extractors.docx import DocxExtractor

    path = tmp_path / "assignment.docx"
    document = Document()
    document.add_heading("Introduction", level=1)
    document.add_paragraph("A formatted paragraph.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First"
    table.cell(0, 1).text = "Second"
    document.save(path)
    result = DocxExtractor().extract(path)
    assert "A formatted paragraph." in result.text
    assert "First | Second" in result.text
    assert result.formatting["heading_count"] == 1
    assert result.formatting["margins"]


def test_txt_extractor_reports_non_utf8_content(tmp_path):
    from app.extractors.text import TextExtractor

    path = tmp_path / "bad.txt"
    path.write_bytes(b"\xff")
    with pytest.raises(ExtractionError, match="UTF-8"):
        TextExtractor().extract(path)


def test_pdf_extracts_text_from_document(tmp_path):
    from reportlab.pdfgen import canvas

    from app.extractors.pdf import PdfExtractor

    path = tmp_path / "readable.pdf"
    document = canvas.Canvas(str(path))
    document.drawString(50, 750, "Text from a PDF document")
    document.save()
    result = PdfExtractor().extract(path)
    assert "Text from a PDF document" in result.text


def test_pdf_corruption_has_friendly_error(tmp_path):
    from app.extractors.pdf import PdfExtractor

    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(ExtractionError, match="corrupted"):
        PdfExtractor().extract(path)


def test_pdf_ocr_fallback_success_and_empty_result(tmp_path, monkeypatch):
    from reportlab.pdfgen import canvas

    from app.extractors.pdf import PdfExtractor

    path = tmp_path / "scanned.pdf"
    document = canvas.Canvas(str(path))
    document.save()
    monkeypatch.setattr(PdfExtractor, "_ocr", staticmethod(lambda _path: __import__("app.extractors.base", fromlist=["ExtractedDocument"]).ExtractedDocument("OCR result")))
    assert PdfExtractor().extract(path).text == "OCR result"
    monkeypatch.setattr(PdfExtractor, "_ocr", staticmethod(lambda _path: (_ for _ in ()).throw(ExtractionError("OCR setup missing"))))
    with pytest.raises(ExtractionError, match="OCR setup missing"):
        PdfExtractor().extract(path)


def test_image_ocr_reads_text_and_rejects_corrupt_image(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace

    from PIL import Image

    from app.extractors.image import ImageExtractor

    path = tmp_path / "scan.png"
    Image.new("RGB", (20, 20), color="white").save(path)
    monkeypatch.delenv("TESSERACT_CMD", raising=False)
    monkeypatch.delenv("TESSERACT_CMD", raising=False)
    monkeypatch.setitem(sys.modules, "pytesseract", SimpleNamespace(image_to_string=lambda _image: "recognized text"))
    assert ImageExtractor().extract(path).text == "recognized text"
    broken = tmp_path / "broken.jpg"
    broken.write_bytes(b"not an image")
    with pytest.raises(ExtractionError, match="corrupted"):
        ImageExtractor().extract(broken)


def test_image_ocr_reports_missing_dependencies_and_empty_recognition(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace

    from PIL import Image

    from app.extractors.image import ImageExtractor

    path = tmp_path / "blank.png"
    Image.new("RGB", (20, 20), color="white").save(path)
    monkeypatch.setitem(sys.modules, "pytesseract", None)
    with pytest.raises(ExtractionError, match="Tesseract"):
        ImageExtractor().extract(path)
    monkeypatch.setitem(sys.modules, "pytesseract", SimpleNamespace(image_to_string=lambda _image: ""))
    with pytest.raises(ExtractionError, match="no readable text"):
        ImageExtractor().extract(path)


def test_pdf_ocr_runs_with_local_dependencies_mocked(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace

    from app.extractors.pdf import PdfExtractor

    monkeypatch.delenv("TESSERACT_CMD", raising=False)
    monkeypatch.delenv("TESSERACT_CMD", raising=False)
    monkeypatch.setitem(
        sys.modules,
        "pdf2image",
        SimpleNamespace(convert_from_path=lambda _path, timeout: ["page-image"]),
    )
    monkeypatch.setitem(
        sys.modules,
        "pytesseract",
        SimpleNamespace(image_to_string=lambda image: f"text from {image}"),
    )
    assert PdfExtractor._ocr(tmp_path / "scan.pdf").text == "text from page-image"


def test_pdf_ocr_missing_tools_is_explained(tmp_path, monkeypatch):
    import sys

    from app.extractors.pdf import PdfExtractor

    monkeypatch.setitem(sys.modules, "pdf2image", None)
    with pytest.raises(ExtractionError, match="Poppler and Tesseract"):
        PdfExtractor._ocr(tmp_path / "scan.pdf")


def test_legacy_ppt_converts_then_extracts_with_libreoffice_mock(tmp_path, monkeypatch):
    from pathlib import Path
    from pptx import Presentation
    from pptx.util import Inches

    from app.extractors.slides import LegacyPowerPointExtractor

    source = tmp_path / "legacy.ppt"
    source.write_bytes(b"mock legacy presentation")
    monkeypatch.setattr("app.extractors.slides.shutil.which", lambda _name: "soffice")

    def convert(arguments, **_kwargs):
        output_dir = Path(arguments[arguments.index("--outdir") + 1])
        output = Presentation()
        slide = output.slides.add_slide(output.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        shape.text = "Converted slide content"
        output.save(output_dir / "legacy.pptx")

    monkeypatch.setattr("app.extractors.slides.subprocess.run", convert)
    result = LegacyPowerPointExtractor().extract(source)
    assert "Converted slide content" in result.text
